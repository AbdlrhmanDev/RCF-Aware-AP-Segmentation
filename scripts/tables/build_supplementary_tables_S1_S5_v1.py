# ================================================================
# PRAD — ALL FIVE SUPPLEMENTARY TABLES S1-S5
# Reads original frozen source CSV files on Google Drive.
# Creates ONE fully populated journal-style DOCX + audit CSVs.
# Does not read radiographs or masks, infer, retrain or modify originals.
# ================================================================
import sys, subprocess, hashlib, json
from pathlib import Path
try:
    from google.colab import drive
    DRIVE = Path('/content/drive/MyDrive')
    if not DRIVE.is_dir():
        drive.mount('/content/drive')
except ImportError:
    DRIVE = Path('/mnt/data/_mock_drive/MyDrive')  # local QA only
try:
    from docx import Document
except ImportError:
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', '-q', 'python-docx'])
    from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pandas as pd
import numpy as np

ROOT = DRIVE / 'RCF_Aware_AP_Segmentation_Research'
PUB = ROOT / '08_Manuscript_Tables_Figures_v1'
TEST = ROOT / '07_Heldout_Test_Evaluation_v1'
VAL = ROOT / '06_Lesion_Level_Analysis'
PLAN = VAL / 'heldout_test_analysis_plan_frozen_v1.json'
EXPECTED_SHA = '35dbf447bbabb1cec4bacebe5ade8c90d9d9fecf4a6e16f40db4ac1add635800'
FILES = {
  'S1': PUB / 'Table_3_Test_PerSeed_PerModel_Results.csv',
  'S2': TEST / 'heldout_test_three_seed_sensitivity_mean_sd_v1.csv',
  'S3': TEST / 'heldout_test_three_seed_validation_vs_test_descriptive_v1.csv',
  'S4': PUB / 'Table_S1_Exploratory_Negative_Subgroups.csv',
  'S5': PUB / 'Table_S2_Exploratory_Paired_FP_Counts.csv',
}
def digest(p):
    h=hashlib.sha256()
    with open(p,'rb') as f:
        for chunk in iter(lambda:f.read(1<<20),b''):h.update(chunk)
    return h.hexdigest()

assert PLAN.is_file(), f'MISSING frozen plan: {PLAN}'
assert digest(PLAN) == EXPECTED_SHA, 'FROZEN PLAN HASH MISMATCH: STOP'
for key,path in FILES.items():
    assert path.is_file(), (f'MISSING {key}: {path}\n'
                             'Locate this exact file in Google Drive and ensure the same account is mounted.')
source_sha = {k:digest(p) for k,p in FILES.items()}
dfs = {k:pd.read_csv(p) for k,p in FILES.items()}

def require(k,cols,rows):
    d=dfs[k]
    missing=set(cols)-set(d.columns)
    assert not missing, f'{k}: Missing columns {sorted(missing)}. Available: {list(d.columns)}'
    assert len(d)==rows, f'{k}: Expected {rows} rows, found {len(d)}'

require('S1',['seed','model','GT','TP','FP','FN','Lesion_Recall',
  'Lesion_Precision','Lesion_F1','FP_Components_per_Image',
  'AP_Negative_FP_Image_Rate','AP_Dice','AP_IoU'],12)
require('S2',['model','IoU_threshold','min_pred_area','N_seeds']+
  [f'{q}_{m}' for q in ['Mean','SD'] for m in [
      'Lesion_F1','Lesion_Recall','Lesion_Precision',
      'FP_Components_per_Image','AP_Negative_FP_Image_Rate']],16)
metrics=['Lesion_F1','Lesion_Recall','Lesion_Precision',
         'FP_Components_per_Image','AP_Negative_FP_Image_Rate']
require('S3',['seed','model']+
  [f'{m}_{split}' for m in metrics for split in ['Validation','Test']]+
  [f'Delta_Test_minus_Validation_{m}' for m in metrics],12)
require('S4',['seed','model','Group','Images','FP',
              'FP_Components_per_Image','Negative_FP_Image_Rate'],12)
require('S5',['seed','Group','Images','M6_total_FP','C2_total_FP',
  'Delta_FP_M6_minus_C2','Images_M6_Fewer_FP',
  'Images_C2_Fewer_FP','Images_FP_Tie'],6)
seeds={42,123,2026}; models={'M5','M6','C1','C2'}
assert set(dfs['S1']['seed'])==seeds and set(dfs['S1']['model'])==models
assert set(dfs['S3']['seed'])==seeds and set(dfs['S3']['model'])==models
assert set(dfs['S2']['model'])==models and set(dfs['S2']['min_pred_area'])=={1,16}
assert set(np.round(dfs['S2']['IoU_threshold'],3))=={0.1,0.5}
assert dfs['S2']['N_seeds'].eq(3).all()
assert set(dfs['S4']['model'])=={'M6','C2'}
assert set(dfs['S4']['Group'])=={'RCF_ONLY','NEITHER'}
assert set(dfs['S5']['Group'])=={'RCF_ONLY','NEITHER'}
for key,keys in [('S1',['seed','model']),('S2',['model','IoU_threshold','min_pred_area']),
                 ('S3',['seed','model']),('S4',['seed','model','Group']),
                 ('S5',['seed','Group'])]:
    assert not dfs[key].duplicated(keys).any(),f'{key}: Duplicate row keys'
# Primary sensitivity values must agree with S1 checkpoint-mean F1 and FP/Image.
primary_s2=dfs['S2'][(dfs['S2']['IoU_threshold'].sub(.1).abs()<1e-10) &
                     (dfs['S2']['min_pred_area']==1)]
assert len(primary_s2)==4
for r in primary_s2.itertuples(index=False):
    s1=dfs['S1'].loc[dfs['S1']['model']==r.model]
    assert np.isclose(s1['Lesion_F1'].mean(),r.Mean_Lesion_F1,atol=1e-8)
    assert np.isclose(s1['FP_Components_per_Image'].mean(),r.Mean_FP_Components_per_Image,atol=1e-8)
for r in dfs['S3'].itertuples(index=False):
    ref=dfs['S1'].loc[(dfs['S1']['model']==r.model)&(dfs['S1']['seed']==r.seed)].iloc[0]
    assert np.isclose(ref['Lesion_F1'],r.Lesion_F1_Test,atol=1e-8)
    assert np.isclose(ref['FP_Components_per_Image'],r.FP_Components_per_Image_Test,atol=1e-8)
    for m in metrics:
        assert np.isclose(getattr(r,'Delta_Test_minus_Validation_'+m),
                          getattr(r,m+'_Test')-getattr(r,m+'_Validation'),atol=1e-8)
for r in dfs['S5'].itertuples(index=False):
    assert r.M6_total_FP-r.C2_total_FP == r.Delta_FP_M6_minus_C2
    assert r.Images_M6_Fewer_FP+r.Images_C2_Fewer_FP+r.Images_FP_Tie==r.Images
    for model,col in [('M6','M6_total_FP'),('C2','C2_total_FP')]:
        sub=dfs['S4'].loc[(dfs['S4']['seed']==r.seed)&
                (dfs['S4']['model']==model)&(dfs['S4']['Group']==r.Group)]
        assert len(sub)==1 and int(sub['FP'].iloc[0])==getattr(r,col)

DEST=ROOT/'11_Supplementary_Tables_Publication_v1'
assert not DEST.exists(), (f'OUTPUT ALREADY EXISTS: {DEST}\n'
                           'Existing outputs protected. Inspect them; do not overwrite.')
DEST.mkdir(parents=True,exist_ok=False)

# ----- Document helpers ------------------------------------------
doc=Document();section=doc.sections[0]
section.orientation=WD_ORIENT.LANDSCAPE
section.page_width,section.page_height=Inches(11.69),Inches(8.27)
section.left_margin=Inches(.58);section.right_margin=Inches(.58)
section.top_margin=Inches(.55);section.bottom_margin=Inches(.55)
styles=doc.styles
styles['Normal'].font.name='Arial';styles['Normal'].font.size=Pt(8.5)
styles['Normal'].paragraph_format.space_after=Pt(4)
for s in ['Heading 1','Heading 2']:
    styles[s].font.name='Arial'
    styles[s].font.color.rgb=RGBColor(24,48,72)
styles['Heading 1'].font.size=Pt(13)
styles['Heading 2'].font.size=Pt(10)

def text(v,prec=4):
    if pd.isna(v):return '—'
    if isinstance(v,(float,np.floating)):
        return f'{float(v):.{prec}f}'
    if isinstance(v,(int,np.integer)):return str(int(v))
    return str(v)

def table(title,data,labels,note=None,prec=4):
    assert len(data.columns)==len(labels)
    doc.add_heading(title,level=2)
    t=doc.add_table(rows=1,cols=len(labels))
    t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for j,label in enumerate(labels):
        cell=t.rows[0].cells[j];cell.text=str(label)
        shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'E6EDF4');cell._tc.get_or_add_tcPr().append(shd)
        for p in cell.paragraphs:
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:r.bold=True;r.font.size=Pt(7.5)
    trPr=t.rows[0]._tr.get_or_add_trPr();header=OxmlElement('w:tblHeader');header.set(qn('w:val'),'true');trPr.append(header)
    for i,row in enumerate(data.itertuples(index=False,name=None)):
        cells=t.add_row().cells
        for j,val in enumerate(row):
            cells[j].text=text(val,prec)
            cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[j].paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after=Pt(0)
                for r in p.runs:r.font.size=Pt(7.0)
        if i%2:
            for cell in cells:
                shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'F6F8FA');cell._tc.get_or_add_tcPr().append(shd)
    if note:
        p=doc.add_paragraph('Note. '+note)
        p.style='Normal'
        for run in p.runs:run.italic=True;run.font.size=Pt(7.6)

def panel(df,cols):
    return df[cols].copy()

p=doc.add_heading('Supplementary Tables S1–S5',0)
p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=doc.add_paragraph('Root Canal Filling–Aware Multitask Segmentation for Apical Periodontitis on Periapical Radiographs')
p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=doc.add_paragraph('Source: frozen internal held-out test, n = 743; three training seeds; 116 reference AP components. No model or threshold retuning.')
p.alignment=WD_ALIGN_PARAGRAPH.CENTER

a=dfs['S1'].sort_values(['seed','model'])
table('Supplementary Table S1. Complete checkpoint-level performance — Panel A: lesion metrics',
      panel(a,['seed','model','GT','TP','FP','FN','Lesion_Recall','Lesion_Precision','Lesion_F1']),
      ['Seed','Model','GT','TP','FP','FN','Recall','Precision','F1'],
      'IoU >= 0.10; minimum predicted component area = 1 pixel; component-level one-to-one matching.')
table('Supplementary Table S1 (continued) — Panel B: additional metrics',
      panel(a,['seed','model','FP_Components_per_Image','AP_Negative_FP_Image_Rate','AP_Dice','AP_IoU']),
      ['Seed','Model','FP / image','AP-negative FP image rate','AP Dice','AP IoU'],
      'AP Dice and IoU are averaged over AP-positive reference images; all models evaluated on the same test partition.')

doc.add_page_break()
s=dfs['S2'].sort_values(['IoU_threshold','min_pred_area','model'])
s2=pd.DataFrame({
    'IoU':s['IoU_threshold'], 'Area':s['min_pred_area'],'Model':s['model'],
    'F1 Mean ± SD':[f'{x:.4f} ± {y:.4f}' for x,y in zip(s.Mean_Lesion_F1,s.SD_Lesion_F1)],
    'Recall Mean ± SD':[f'{x:.4f} ± {y:.4f}' for x,y in zip(s.Mean_Lesion_Recall,s.SD_Lesion_Recall)],
    'Precision Mean ± SD':[f'{x:.4f} ± {y:.4f}' for x,y in zip(s.Mean_Lesion_Precision,s.SD_Lesion_Precision)],
    'FP/image Mean ± SD':[f'{x:.4f} ± {y:.4f}' for x,y in zip(s.Mean_FP_Components_per_Image,s.SD_FP_Components_per_Image)],
    'Neg-FP rate Mean ± SD':[f'{x:.4f} ± {y:.4f}' for x,y in zip(s.Mean_AP_Negative_FP_Image_Rate,s.SD_AP_Negative_FP_Image_Rate)],
})
table('Supplementary Table S2. Sensitivity to component-matching IoU and minimum predicted area',
      s2,['IoU','Area (px)','Model','F1 (Mean ± SD)','Recall (Mean ± SD)','Precision (Mean ± SD)','FP/image (Mean ± SD)','Neg-FP rate (Mean ± SD)'],
      'Values summarize the same three training seeds; IoU 0.10 and area 1 pixel are the frozen primary setting. These are descriptive means and sample SDs, not bootstrap confidence intervals.')

doc.add_page_break()
v=dfs['S3'].sort_values(['seed','model'])
s3a=pd.DataFrame({'Seed':v.seed,'Model':v.model,
  'F1 Validation':v.Lesion_F1_Validation,'F1 Test':v.Lesion_F1_Test,
  'Δ F1 Test−Val':v.Delta_Test_minus_Validation_Lesion_F1,
  'Recall Validation':v.Lesion_Recall_Validation,'Recall Test':v.Lesion_Recall_Test,
  'Δ Recall Test−Val':v.Delta_Test_minus_Validation_Lesion_Recall})
table('Supplementary Table S3. Validation-to-test comparison — Panel A: lesion F1 and recall',
      s3a,list(s3a.columns),
      'Descriptive only: validation was used for checkpoint selection. Δ = Test minus Validation.')
s3b=pd.DataFrame({'Seed':v.seed,'Model':v.model,
 'Precision Validation':v.Lesion_Precision_Validation,
 'Precision Test':v.Lesion_Precision_Test,
 'Δ Precision Test−Val':v.Delta_Test_minus_Validation_Lesion_Precision,
 'FP/image Validation':v.FP_Components_per_Image_Validation,
 'FP/image Test':v.FP_Components_per_Image_Test,
 'Δ FP/image Test−Val':v.Delta_Test_minus_Validation_FP_Components_per_Image})
table('Supplementary Table S3 (continued) — Panel B: lesion precision and false-positive components',
      s3b,list(s3b.columns),
      'Unmatched AP components divided by all evaluated images. Δ = Test minus Validation.')
s3c=pd.DataFrame({'Seed':v.seed,'Model':v.model,
 'Neg-FP rate Validation':v.AP_Negative_FP_Image_Rate_Validation,
 'Neg-FP rate Test':v.AP_Negative_FP_Image_Rate_Test,
 'Δ Neg-FP rate Test−Val':v.Delta_Test_minus_Validation_AP_Negative_FP_Image_Rate})
table('Supplementary Table S3 (continued) — Panel C: AP-negative false-positive image rate',
      s3c,list(s3c.columns),
      'Fraction of reference AP-negative images with any positive AP prediction. Δ = Test minus Validation.')

doc.add_page_break()
b=dfs['S4'].sort_values(['seed','Group','model'])
table('Supplementary Table S4. Exploratory AP-negative subgroup false-positive burden',
      panel(b,['seed','model','Group','Images','FP','FP_Components_per_Image','Negative_FP_Image_Rate']),
      ['Seed','Model','Group','Images','FP components','FP/image','Neg-FP image rate'],
      'Post-hoc descriptive analysis; RCF_ONLY n=245 and NEITHER n=399. These groups have no annotated AP.')

c=dfs['S5'].sort_values(['seed','Group'])
table('Supplementary Table S5. Exploratory paired false-positive differences, M6 vs C2',
      panel(c,['seed','Group','Images','M6_total_FP','C2_total_FP',
       'Delta_FP_M6_minus_C2','Images_M6_Fewer_FP','Images_C2_Fewer_FP','Images_FP_Tie']),
      ['Seed','Group','Images','M6 FP','C2 FP','Δ FP (M6−C2)',
       'Images M6 fewer FP','Images C2 fewer FP','Ties'],
      'Post-hoc descriptive analysis. Negative Δ indicates fewer unmatched components for M6. Counts of images with fewer FP are paired image comparisons.')

out=DEST/'PRAD_Supplementary_Tables_S1-S5_Complete_v1.docx'
doc.save(out)
for key,frame in dfs.items():
    frame.to_csv(DEST/f'Supplementary_{key}_Full_Precision_v1.csv',index=False)
metadata={
 'frozen_plan_sha256':EXPECTED_SHA,
 'source_files':{k:{'path':str(FILES[k]),'sha256':source_sha[k],'rows':len(dfs[k])}
                 for k in FILES},
 'original_test_results_modified':False,
 'model_training_or_inference':False,
 'table_rounding': '4 decimal places in DOCX; full precision retained in CSV',
 'post_hoc_tables':['S4','S5'], 'patient_disjointness_verified':False,
}
with open(DEST/'supplementary_table_provenance_v1.json','w',encoding='utf-8') as f:
    json.dump(metadata,f,ensure_ascii=False,indent=2)
assert out.is_file() and out.stat().st_size>10000
print('='*85)
print('ALL FIVE SUPPLEMENTARY TABLES S1–S5 — COMPLETE AND AUDITED')
print('='*85)
for k in FILES:print(f'{k}: {len(dfs[k])} source rows; SHA {source_sha[k][:16]}')
print('Word document:',out)
print('Original input files: NOT MODIFIED')
print('Model inference or training: NONE')
print('Frozen test plan: VERIFIED')
print('FINAL STATUS: PASSED')
